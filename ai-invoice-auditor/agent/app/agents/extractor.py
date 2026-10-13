"""Extractor Agent (no LLM).

1. Raw text:  PDF -> pdfplumber, OCR fallback if no text layer | DOCX -> python-docx | PNG -> Tesseract
2. Fields:    rule-based, multilingual (en/es/de) parsing into header + line items.

Missing values are returned as None (never guessed) so the validation agents can flag them.
"""
import logging
import os
import re
from datetime import date
from pathlib import Path

import pdfplumber
import pytesseract
import yaml
from docx import Document
from pdf2image import convert_from_path
from PIL import Image, ImageOps

log = logging.getLogger(__name__)

OCR_LANGS = os.environ.get("OCR_LANGS", "eng+spa+deu")   # must match packs installed in agent/Dockerfile
MIN_TEXT_CHARS = 20                                       # below this a PDF is treated as a scan
DEFAULT_SYMBOLS = {"$": "USD", "€": "EUR", "₹": "INR", "£": "GBP"}
ISO_CODES = ("USD", "EUR", "INR", "GBP")


# --------------------------------------------------------------------------- raw text
OCR_PASSES = [(1, 4), (2, 6), (3, 4), (3, 6), (2, 4)]    # (upscale factor, tesseract psm)


def _prep(img: Image.Image) -> Image.Image:
    if img.mode in ("RGBA", "LA", "P"):                    # flatten transparency onto white
        img = img.convert("RGBA")
        bg = Image.new("RGBA", img.size, "white")
        img = Image.alpha_composite(bg, img)
    return ImageOps.autocontrast(ImageOps.grayscale(img.convert("RGB")))


def _score(text: str) -> float:
    """How plausible is this OCR result? Reward filled fields and numbers that add up."""
    h, items = parse_header(text, ocr=True), parse_line_items(text)
    sc = sum(1 for k in ("invoice_no", "invoice_date", "vendor_name", "po_number", "total_amount") if h.get(k))
    if h["subtotal"] is not None and h["tax_amount"] is not None and h["total_amount"] is not None:
        sc += 3 if abs(h["subtotal"] + h["tax_amount"] - h["total_amount"]) < 0.01 else 0
    if items and h["subtotal"] is not None:
        sc += 2 if abs(sum(i["total"] or 0 for i in items) - h["subtotal"]) < 0.01 else 0
    sc += 1 if items else 0
    if h.get("invoice_no") and re.fullmatch(r"[A-Z]{2,4}-\d+", h["invoice_no"]):
        sc += 1
    return sc


def _ocr(img: Image.Image) -> str:
    """Several scale/psm passes; keep the text whose parsed numbers are most self-consistent."""
    base, best, best_score = _prep(img), "", -1.0
    for scale, psm in OCR_PASSES:
        im = base if scale == 1 else base.resize((base.width * scale, base.height * scale), Image.LANCZOS)
        text = pytesseract.image_to_string(im, lang=OCR_LANGS, config=f"--psm {psm}")
        sc = _score(text)
        if sc > best_score:
            best, best_score = text, sc
    return best


def extract_pdf(path: Path) -> tuple[str, str]:
    with pdfplumber.open(path) as pdf:
        text = "\n".join((page.extract_text() or "") for page in pdf.pages).strip()
    if len(text) >= MIN_TEXT_CHARS:
        return text, "pdfplumber"
    log.info("%s has no text layer -> OCR", path.name)
    return "\n".join(_ocr(p) for p in convert_from_path(str(path), dpi=300)).strip(), "tesseract"


def extract_docx(path: Path) -> tuple[str, str]:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        parts += [" | ".join(c.text.strip() for c in row.cells) for row in t.rows]
    return "\n".join(parts).strip(), "docx"


def extract_image(path: Path) -> tuple[str, str]:
    return _ocr(Image.open(path)).strip(), "tesseract"


ROUTES = {".pdf": extract_pdf, ".docx": extract_docx, ".png": extract_image}


def extract_text(path: str) -> tuple[str, str]:
    p = Path(path)
    try:
        return ROUTES[p.suffix.lower()](p)
    except KeyError:
        raise ValueError(f"Unsupported file type: {p.suffix}")


# --------------------------------------------------------------------------- helpers
def _symbol_map() -> dict:
    try:
        return yaml.safe_load(Path(os.environ["RULES_PATH"]).read_text())["currency_symbol_map"]
    except Exception:
        return DEFAULT_SYMBOLS


NUM = r"\d[\d.,]*"


def parse_number(s: str | None) -> float | None:
    """'1,234.56' / '1.234,56' / '2,50' / '$1470.00' / '40000' -> float."""
    if s is None:
        return None
    if isinstance(s, (int, float)):
        return round(float(s), 4)
    s = re.sub(r"[^\d.,\-]", "", s.replace(" ", ""))
    if not re.search(r"\d", s):
        return None
    if "," in s and "." in s:
        dec = "," if s.rfind(",") > s.rfind(".") else "."
        s = s.replace("." if dec == "," else ",", "").replace(dec, ".")
    elif "," in s:
        s = s.replace(",", ".") if re.search(r",\d{1,2}$", s) else s.replace(",", "")
    try:
        return round(float(s), 4)
    except ValueError:
        return None


_MONTHS = {}
for _names in (
    "january february march april may june july august september october november december",
    "enero febrero marzo abril mayo junio julio agosto septiembre octubre noviembre diciembre",
    "januar februar märz april mai juni juli august september oktober november dezember",
):
    for _i, _n in enumerate(_names.split(), 1):
        _MONTHS[_n] = _i
        _MONTHS[_n[:3]] = _i
_MONTHS.update({"maerz": 3, "sept": 9, "setiembre": 9})


def parse_date(s: str | None) -> str | None:
    """Return ISO date string or None. Handles 2025-03-14, 02.05.2025, 14 March 2025, 10 de abril de 2025, 25-May-2025."""
    if not s:
        return None
    s = s.strip()
    try:
        if m := re.search(r"(\d{4})-(\d{1,2})-(\d{1,2})", s):
            return date(int(m[1]), int(m[2]), int(m[3])).isoformat()
        if m := re.search(r"(\d{1,2})[./](\d{1,2})[./](\d{4})", s):          # day-first (es/de)
            return date(int(m[3]), int(m[2]), int(m[1])).isoformat()
        if m := re.search(r"(\d{1,2})[\s\-]+(?:de\s+)?([^\W\d_]+)\.?[\s\-,]+(?:de\s+)?(\d{4})", s):
            mon = _MONTHS.get(m[2].lower())
            if mon:
                return date(int(m[3]), mon, int(m[1])).isoformat()
    except ValueError:
        pass
    return None


def detect_currency(text: str) -> str | None:
    for sym, iso in _symbol_map().items():            # symbol must touch a number (ignores OCR noise like "$F Lane")
        e = re.escape(sym)
        if re.search(rf"{e}\s*\d|\d\s*{e}", text):
            return iso
    if m := re.search(r"\b(" + "|".join(ISO_CODES) + r")\b", text):
        return m[1]
    return None


def _last_number(line: str) -> float | None:
    nums = re.findall(r"\d[\d.,]*", line)
    return parse_number(nums[-1]) if nums else None


def _labelled_amount(text: str, labels: str) -> float | None:
    """Last matching line that STARTS with one of the labels (so 'Subtotal' never matches 'Total')."""
    hit = None
    for m in re.finditer(rf"(?im)^\s*(?:{labels})\b[^\n]*$", text):
        v = _last_number(m[0])
        if v is not None:
            hit = v
    return hit


# --------------------------------------------------------------------------- header
LABEL_INVOICE_NO = (r"invoice(?:\s*(?:no\.?|number|num|#))?|factura(?:\s*(?:no\.?|n[úu]mero|n[º°]))?|"
                    r"n[úu]mero de factura|rechnungs?(?:nummer|\s*nr\.?)|facture")
LABEL_DATE = r"invoice\s+date|date|fecha(?:\s+de\s+factura)?|rechnungsdatum|datum"
LABEL_TOTAL = (r"total\s+amount\s+due|total\s+due|amount\s+due|grand\s+total|total\s+a\s+pagar|importe\s+total|"
               r"gesamtbetrag|gesamtsumme|total")
LABEL_SUBTOTAL = r"subtotal|zwischensumme|base\s+imponible"
LABEL_TAX = r"tax|vat|gst|iva|impuesto|mwst|ust|steuer"


_OCR_DIGIT = str.maketrans({"L": "1", "l": "1", "I": "1", "|": "1", "O": "0", "o": "0", "S": "5", "B": "8", "Z": "2"})


def fix_ocr_number(token: str) -> str:
    """INV-L00S -> INV-1005 : in the numeric part, map look-alike letters to digits."""
    prefix, sep, rest = token.partition("-")
    if not sep or not rest:
        return token
    fixed = rest.translate(_OCR_DIGIT)
    return f"{prefix}-{fixed}" if fixed.isdigit() else token


def parse_header(text: str, ocr: bool = False) -> dict:
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    vendor = lines[0] if lines else None

    m = re.search(rf"(?im)^\s*(?:{LABEL_INVOICE_NO})\s*[:#]?\s*(?=[A-Za-z0-9\-/]*\d)([A-Za-z0-9\-/]+)", text)
    invoice_no = m[1] if m else None
    if ocr:                                                     # scans: garbled label / look-alike characters
        if not invoice_no:
            for c in re.finditer(r"\b([A-Z]{2,4}-[A-Z0-9]{3,})\b", text):
                if not c[1].startswith(("PO-", "SKU-")):
                    invoice_no = c[1]
                    break
        if invoice_no:
            invoice_no = fix_ocr_number(invoice_no)

    m = re.search(rf"(?im)^\s*(?:{LABEL_DATE})\s*:?\s*(.+)$", text)
    invoice_date = parse_date(m[1]) if m else None

    m = re.search(r"\bPO[-\s]?\d+\b", text, re.I)
    po_number = re.sub(r"\s", "-", m[0]).upper().replace("PO-", "PO-") if m else None
    if po_number and not po_number.startswith("PO-"):
        po_number = "PO-" + po_number[2:]

    tax_rate = None
    if m := re.search(rf"(?im)^\s*(?:{LABEL_TAX})\b[^\n]*?\(?\s*(\d+(?:[.,]\d+)?)\s*%", text):
        tax_rate = parse_number(m[1])

    return {
        "invoice_no": invoice_no,
        "invoice_date": invoice_date,
        "vendor_name": vendor,
        "vendor_id": None,                       # resolved from vendor_name via mock-erp (business validation)
        "po_number": po_number,
        "currency": detect_currency(text),
        "subtotal": _labelled_amount(text, LABEL_SUBTOTAL),
        "tax_amount": _labelled_amount(text, LABEL_TAX),
        "tax_rate": tax_rate,
        "total_amount": _labelled_amount(text, LABEL_TOTAL),
    }


# --------------------------------------------------------------------------- line items
CODE_RE = re.compile(r"^[A-Za-z]{2,}-?\d+$")
STOP_RE = re.compile(rf"(?i)^\s*(?:{LABEL_SUBTOTAL}|{LABEL_TOTAL}|{LABEL_TAX})\b")


def _item(n, code, desc, qty, unit, total) -> dict:
    qty, unit, total = parse_number(qty), parse_number(unit), parse_number(total)
    if total is None and qty is not None and unit is not None:
        total = round(qty * unit, 2)
    return {"line_number": n, "item_code": code or None, "description": (desc or "").strip() or None,
            "qty": qty, "unit_price": unit, "total": total}


def parse_line_items(text: str) -> list[dict]:
    items: list[dict] = []

    # A) pipe tables:  CODE | description | qty | unit | total   (pdf / docx / es / de)
    for line in text.splitlines():
        cells = [c.strip() for c in line.split("|")]
        if len(cells) >= 5 and CODE_RE.match(cells[0]):
            items.append(_item(len(items) + 1, cells[0], cells[1], cells[2], cells[3], cells[4]))
    if items:
        return items

    # B) key/value single item:  Description: X / Quantity: 200 units / Unit Price: $1.25
    d = re.search(r"(?im)^\s*(?:description|descripci[óo]n|beschreibung)\s*:\s*(.+)$", text)
    q = re.search(rf"(?im)^\s*(?:quantity|qty|cantidad|menge)\s*:\s*({NUM})", text)
    u = re.search(rf"(?im)^\s*(?:unit\s*price|precio(?:\s*unitario)?|preis|einzelpreis)\s*:\s*[^\d\n]*({NUM})", text)
    if d and q and u:
        code = re.search(r"\bSKU-\d+\b", text)
        total = _labelled_amount(text, LABEL_SUBTOTAL)
        return [_item(1, code[0] if code else None, d[1], q[1], u[1], total)]

    # C) whitespace/OCR rows without codes:  Fuel Surcharge   1  120  120
    row = re.compile(rf"^(?P<d>[^\W\d_][^\n]*?)\s+(?P<q>\d+(?:[.,]\d+)?)\s+(?P<u>\d[\d.,]*)\s+(?P<t>\d[\d.,]*)\s*$")
    started = False
    for line in text.splitlines():
        if re.search(r"(?i)\b(qty|ty|cantidad|menge|unit|umit|precio|preis)\b.*\b(total|gesamt)\b", line):
            started = True                                    # header row seen
            continue
        if started and STOP_RE.match(line):
            break
        if started and (m := row.match(line.strip())):
            items.append(_item(len(items) + 1, None, m["d"], m["q"], m["u"], m["t"]))
    return items


def parse_fields(raw_text: str, ocr: bool = False) -> tuple[dict, list[dict]]:
    return parse_header(raw_text, ocr=ocr), parse_line_items(raw_text)


def extract_invoice(path: str) -> dict:
    """Full extractor output for one file."""
    text, method = extract_text(path)
    header, items = parse_fields(text, ocr=(method == "tesseract"))
    return {"raw_text": text, "extraction_method": method, "header": header, "line_items": items}